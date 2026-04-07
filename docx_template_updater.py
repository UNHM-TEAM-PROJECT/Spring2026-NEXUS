"""Minimal DOCX template updater used by api_routes.

Replaces placeholders like {{field_name}} in paragraphs/tables/headers/footers.
"""

import json
import re
from pathlib import Path
from typing import Any, Dict

from docx import Document


MISSING = "Missing"


def _clean(value: Any) -> str:
    if value is None:
        return MISSING
    if isinstance(value, (dict, list, tuple, set)):
        return MISSING
    text = str(value).strip()
    return text if text else MISSING


def _get(data: Dict[str, Any], *keys: str) -> Any:
    cur: Any = data
    for k in keys:
        if not isinstance(cur, dict):
            return None
        cur = cur.get(k)
    return cur


def _format_structured_text(value: str) -> str:
    """Normalize detector prose into cleaner bullet-style lines."""
    if value == MISSING:
        return value

    raw = (value or "").replace("\r\n", "\n")
    raw = re.sub(r"\s+", " ", raw)
    raw = re.sub(r"\s*\n\s*", "\n", raw)

    split_pattern = r"(?:\n+|\s*;\s*|\s*\|\s*|\s*•\s*|\s*\u2022\s*|\s+-\s+)"
    parts = [p.strip(" .") for p in re.split(split_pattern, raw) if p.strip()]

    # If no clear separators were found, try splitting numbered clauses.
    if len(parts) <= 1:
        numbered = re.split(r"\s+(?=\d+[\.)]\s+)", raw.strip())
        parts = [p.strip(" .") for p in numbered if p.strip()]

    # De-duplicate while preserving order.
    seen = set()
    unique_parts = []
    for part in parts:
        key = part.lower()
        if key in seen:
            continue
        seen.add(key)
        unique_parts.append(part)

    if not unique_parts:
        return value

    if len(unique_parts) == 1:
        return unique_parts[0]

    return "\n".join(f"- {item}" for item in unique_parts)


def _strip_redundant_header(value: str, aliases: list[str]) -> str:
    """Remove detector headers when template already has section headers."""
    if value == MISSING:
        return value

    text = (value or "").strip()
    if not text:
        return value

    alias_pattern = "|".join(re.escape(alias) for alias in aliases if alias)
    if not alias_pattern:
        return value

    # Remove leading markdown/hash bullets and one or more duplicated headers.
    text = re.sub(
        rf"^(?:\s*[#*-]+\s*)?(?:{alias_pattern})\s*[:\-]?\s*",
        "",
        text,
        flags=re.IGNORECASE,
    )

    # Remove standalone header lines that may still exist after newline splits.
    lines = []
    for line in text.splitlines():
        clean = line.strip()
        if re.fullmatch(
            rf"(?:\s*[#*-]+\s*)?(?:{alias_pattern})\s*[:\-]?\s*",
            clean,
            flags=re.IGNORECASE,
        ):
            continue
        lines.append(line)

    cleaned = "\n".join(lines).strip()
    return cleaned if cleaned else value


def _to_template_values(record: Dict[str, Any]) -> Dict[str, str]:
    """Accept either already-flat template keys or nested detector payload."""
    course_code = _clean(record.get("course_code"))
    course_name = _clean(record.get("course_name"))
    course_title = _clean(record.get("course_title"))
    if course_title == MISSING and course_code != MISSING:
        course_title = course_code

    values = {
        "filename": _clean(record.get("filename")),
        "course_title": course_title,
        "course_code": course_code,
        "course_name": course_name,
        "instructor_name": _clean(record.get("instructor_name")),
        "instructor_title": _clean(record.get("instructor_title")),
        "instructor_department": _clean(record.get("instructor_department")),
        "email": _clean(record.get("email")),
        "preferred_contact_method": _clean(
            record.get("preferred_contact_method")
        ),
        "response_time": _clean(record.get("response_time")),
        "office_address": _clean(record.get("office_address")),
        "office_phone": _clean(record.get("office_phone")),
        "office_hours": _clean(record.get("office_hours")),
        "modality": _clean(record.get("modality")),
        "class_location": _clean(record.get("class_location")),
        "assignment_types_title": _clean(
            record.get("assignment_types_title")
        ),
        "assignment_delivery": _clean(record.get("assignment_delivery")),
        "deadline_expectations_title": _clean(
            record.get("deadline_expectations_title")
        ),
        "SLOs": _clean(record.get("SLOs")),
        "credit_hour": _clean(record.get("credit_hour")),
        "workload": _clean(record.get("workload")),
        "grading_process": _clean(record.get("grading_process")),
        "final_grade_scale": _clean(record.get("final_grade_scale")),
    }

    # If flat values are missing, attempt nested detector keys.
    fallback = {
        "instructor_name": _clean(_get(record, "instructor", "name")),
        "instructor_title": _clean(_get(record, "instructor", "title")),
        "instructor_department": _clean(
            _get(record, "instructor", "department")
        ),
        "email": _clean(_get(record, "email_information", "email")),
        "preferred_contact_method": _clean(
            _get(record, "preferred_information", "preferred")
        ),
        "response_time": _clean(_get(record, "response_time", "content")),
        "office_address": _clean(
            _get(record, "office_information", "location")
        ),
        "office_phone": _clean(_get(record, "office_information", "phone")),
        "office_hours": _clean(_get(record, "office_information", "hours")),
        "modality": _clean(record.get("course_delivery")),
        "class_location": _clean(_get(record, "class_location", "content")),
        "assignment_types_title": _clean(
            _get(record, "assignment_types", "content")
        ),
        "assignment_delivery": _clean(
            _get(record, "assignment_delivery", "content")
        ),
        "deadline_expectations_title": _clean(
            _get(record, "late_information", "late")
        ),
        "SLOs": _clean(
            record.get("slo_content") if record.get("has_slos") else None
        ),
        "credit_hour": _clean(_get(record, "credit_hours", "hours")),
        "workload": _clean(
            _get(record, "workload_information", "description")
        ),
        "grading_process": _clean(_get(record, "grading_process", "content")),
        "final_grade_scale": _clean(
            _get(record, "grading_scale", "content")
        ),
    }

    values["grading_scale"] = values["final_grade_scale"]
    values["late_work_policy"] = values["deadline_expectations_title"]
    values["office_location"] = values["office_address"]
    values["preferred_contact"] = values["preferred_contact_method"]

    for k, v in fallback.items():
        if values.get(k) == MISSING and v != MISSING:
            values[k] = v

    header_aliases = {
        "SLOs": [
            "SLOs",
            "SLO",
            "Student Learning Outcomes",
            "Learning Outcomes",
        ],
        "modality": ["Modality", "Course Delivery", "Delivery Mode"],
        "assignment_types_title": ["Assignment Types", "Assignments"],
        "assignment_delivery": [
            "Assignment Delivery",
            "Submission",
            "Deliverables",
        ],
        "deadline_expectations_title": [
            "Late Work",
            "Deadline Expectations",
            "Missing Work",
        ],
        "response_time": ["Response Time", "Turnaround Time"],
        "class_location": ["Class Location", "Location"],
        "workload": ["Workload", "Expected Workload", "Course Workload"],
        "grading_process": [
            "Grading Process",
            "Grading Procedures",
            "Evaluation",
        ],
        "final_grade_scale": [
            "Grading Scale",
            "Final Grade Scale",
            "Grade Scale",
        ],
        "office_address": ["Office", "Office Location", "Location"],
        "office_hours": ["Office Hours", "Hours"],
        "office_phone": ["Phone", "Office Phone", "Contact"],
        "preferred_contact_method": [
            "Preferred Contact",
            "Preferred Contact Method",
            "Contact Method",
        ],
    }

    format_fields = [
        "SLOs",
        "modality",
        "assignment_types_title",
        "assignment_delivery",
        "deadline_expectations_title",
        "response_time",
        "class_location",
        "workload",
        "grading_process",
        "final_grade_scale",
        "office_address",
        "office_hours",
        "preferred_contact_method",
    ]

    for field, aliases in header_aliases.items():
        values[field] = _strip_redundant_header(values[field], aliases)

    for field in format_fields:
        values[field] = _format_structured_text(values[field])

    return values


def _replace_span_in_runs(
    runs,
    start: int,
    end: int,
    replacement: str,
) -> None:
    """Replace a character span [start, end) in paragraph runs."""
    if start >= end:
        return

    # Build char->run index map for current run state.
    idx_map = []
    for run_idx, run in enumerate(runs):
        for char_idx, _ in enumerate(run.text or ""):
            idx_map.append((run_idx, char_idx))

    if not idx_map or start < 0 or end > len(idx_map):
        return

    first_run_idx, first_char_idx = idx_map[start]
    last_run_idx, last_char_idx = idx_map[end - 1]

    first_text = runs[first_run_idx].text or ""
    last_text = runs[last_run_idx].text or ""

    if first_run_idx == last_run_idx:
        runs[first_run_idx].text = (
            first_text[:first_char_idx]
            + replacement
            + first_text[last_char_idx + 1:]
        )
        return

    runs[first_run_idx].text = first_text[:first_char_idx] + replacement

    for i in range(first_run_idx + 1, last_run_idx):
        runs[i].text = ""

    runs[last_run_idx].text = last_text[last_char_idx + 1:]


def _replace_placeholders_in_paragraph(
    paragraph,
    values: Dict[str, str],
) -> None:
    full_text = paragraph.text or ""
    if "{{" not in full_text:
        return

    matches = list(re.finditer(r"\{\{[^{}]+\}\}", full_text))
    if not matches:
        return

    # Replace from end so earlier match indices remain valid.
    for match in reversed(matches):
        token = match.group(0)
        key = token[2:-2].strip()
        replacement = values.get(key)
        if replacement is None:
            continue
        _replace_span_in_runs(
            paragraph.runs,
            match.start(),
            match.end(),
            replacement,
        )


def _replace_all(doc: Document, values: Dict[str, str]) -> None:
    for p in doc.paragraphs:
        _replace_placeholders_in_paragraph(p, values)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    _replace_placeholders_in_paragraph(p, values)

    for section in doc.sections:
        for p in section.header.paragraphs:
            _replace_placeholders_in_paragraph(p, values)
        for p in section.footer.paragraphs:
            _replace_placeholders_in_paragraph(p, values)


def _replace_unresolved(doc: Document) -> None:
    pattern = re.compile(r"\{\{[^{}]+\}\}")

    def replace_unresolved_in_paragraph(paragraph) -> None:
        text = paragraph.text or ""
        if not pattern.search(text):
            return

        unresolved = {
            match.group(0)[2:-2].strip(): MISSING
            for match in pattern.finditer(text)
        }
        _replace_placeholders_in_paragraph(paragraph, unresolved)

    for p in doc.paragraphs:
        replace_unresolved_in_paragraph(p)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    replace_unresolved_in_paragraph(p)

    for section in doc.sections:
        for p in section.header.paragraphs:
            replace_unresolved_in_paragraph(p)
        for p in section.footer.paragraphs:
            replace_unresolved_in_paragraph(p)


def fill_docx_template_from_detector_result(
    template_path: Path,
    detector_result: Dict[str, Any],
    output_path: Path,
) -> None:
    if not isinstance(detector_result, dict):
        raise ValueError("detector_result must be a dict")

    record = detector_result
    if isinstance(record.get("results"), list) and record["results"]:
        record = record["results"][0]

    values = _to_template_values(record)
    doc = Document(str(template_path))
    _replace_all(doc, values)
    _replace_unresolved(doc)
    doc.save(str(output_path))


def fill_docx_template(
    template_path: Path,
    data_path: Path,
    output_path: Path,
    index: int,
) -> None:
    """Small file-based helper kept for compatibility."""
    raw = json.loads(data_path.read_text(encoding="utf-8"))
    if isinstance(raw, dict) and isinstance(raw.get("results"), list):
        raw = raw["results"]
    if isinstance(raw, list):
        if index < 0 or index >= len(raw):
            raise IndexError("index out of range")
        record = raw[index]
    elif isinstance(raw, dict):
        record = raw
    else:
        raise ValueError("JSON must be object or list")

    fill_docx_template_from_detector_result(template_path, record, output_path)
